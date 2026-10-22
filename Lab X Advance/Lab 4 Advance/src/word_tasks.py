# agentic-document-automator/src/word_tasks.py
import os

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .utils import setup_logging

logger = setup_logging(__name__)

TABLE_PLACEHOLDER = "{{TABLE_PLACEHOLDER}}"

class WordTasks:
    """Handles Word document operations (template filling, tables, text extraction)."""

    def __init__(self):
        logger.info("WordTasks initialized.")

    @staticmethod
    def _replace_placeholders_in_paragraph(paragraph, placeholders):
        """
        Replaces every '{{KEY}}' in a paragraph. Word often splits text into several runs,
        so if a placeholder is not inside a single run, the paragraph text is rebuilt in the first run.
        """
        if "{{" not in paragraph.text:
            return
        for key, value in placeholders.items():
            placeholder = f"{{{{{key}}}}}" # -> '{{KEY}}'
            if placeholder not in paragraph.text:
                continue
            replaced_in_run = False
            for run in paragraph.runs:
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, str(value))
                    replaced_in_run = True
            if not replaced_in_run and paragraph.runs:
                full_text = paragraph.text.replace(placeholder, str(value))
                paragraph.runs[0].text = full_text
                for run in paragraph.runs[1:]:
                    run.text = ""

    def _iter_all_paragraphs(self, document):
        """Yields paragraphs in the body, inside tables, and in headers/footers."""
        yield from document.paragraphs
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        for section in document.sections:
            yield from section.header.paragraphs
            yield from section.footer.paragraphs

    @staticmethod
    def _build_table(document, table_data):
        """Creates a table from a list of rows (first row = header)."""
        cols = len(table_data[0])
        table = document.add_table(rows=0, cols=cols)
        try:
            table.style = 'Table Grid'
        except Exception:
            logger.warning("Style 'Table Grid' not available, using default table style.")
        for r_idx, row_data in enumerate(table_data):
            cells = table.add_row().cells
            for c_idx, cell_data in enumerate(row_data[:cols]):
                cells[c_idx].text = str(cell_data)
                if r_idx == 0:
                    for run in cells[c_idx].paragraphs[0].runs:
                        run.bold = True
                    WordTasks._shade_cell(cells[c_idx], "D9E2F3") # light blue header
        return table

    @staticmethod
    def _shade_cell(cell, hex_color):
        """Sets a cell's background colour (python-docx has no direct API for this)."""
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), hex_color)
        cell._tc.get_or_add_tcPr().append(shading)

    def generate_report_from_template(self, template_path, output_path, data_placeholders, table_data=None):
        """
        Generates a Word document from a template.
        `data_placeholders`: {"REPORT_TITLE": "..."} fills '{{REPORT_TITLE}}' in the template.
        `table_data`: list of rows (first row = header), inserted where '{{TABLE_PLACEHOLDER}}' is.
        """
        if not os.path.exists(template_path):
            logger.error(f"Word template not found: {template_path}")
            return False

        try:
            document = docx.Document(template_path)

            if table_data:
                table = self._build_table(document, table_data)
                anchor = next((p for p in document.paragraphs if TABLE_PLACEHOLDER in p.text), None)
                if anchor is not None:
                    # add_table() appends at the end of the body; move it to where the placeholder was
                    anchor._element.addnext(table._tbl)
                    anchor._element.getparent().remove(anchor._element)
                    logger.info(f"Inserted {len(table_data)}-row table at {TABLE_PLACEHOLDER}.")
                else:
                    logger.warning(f"{TABLE_PLACEHOLDER} not found in template. Table appended to the end.")

            for paragraph in self._iter_all_paragraphs(document):
                self._replace_placeholders_in_paragraph(paragraph, data_placeholders)

            leftover = sorted({p.text for p in self._iter_all_paragraphs(document) if "{{" in p.text})
            if leftover:
                logger.warning(f"Placeholders without data remain in the report: {leftover}")

            os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
            document.save(output_path)
            logger.info(f"Report generated from template and saved to '{output_path}'.")
            return True

        except Exception as e:
            logger.error(f"Error generating report from template: {e}")
            return False

    def extract_text_from_word(self, docx_path):
        """Extracts all text from a Word document (paragraphs, then tables as tab-separated rows)."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None
        try:
            document = docx.Document(docx_path)
            full_text = [para.text for para in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    full_text.append("\t".join(cell.text for cell in row.cells))
            logger.info(f"Text extracted from Word document: {docx_path}")
            return "\n".join(full_text)
        except Exception as e:
            logger.error(f"Error extracting text from Word document '{docx_path}': {e}")
            return None
