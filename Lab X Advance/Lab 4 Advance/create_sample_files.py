# agentic-document-automator/create_sample_files.py
# Creates the input files the agent needs:
#   templates/report_template.docx  - Word template with {{PLACEHOLDER}} fields
#   documents/source_document.pdf   - 2-page PDF to be watermarked / text-extracted
#   documents/dummy_doc1.pdf        - Part A of the project documentation (for merging)
#   documents/dummy_doc2.pdf        - Part B of the project documentation (for merging)
import os

import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
DOCUMENTS_DIR = os.path.join(BASE_DIR, "documents")


def create_word_template(path):
    document = docx.Document()

    title = document.add_heading("{{REPORT_TITLE}}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for label, placeholder in [("Date", "{{REPORT_DATE}}"), ("Author", "{{AUTHOR_NAME}}"),
                               ("Company", "{{COMPANY_NAME}}"), ("Period", "{{MONTH_PERIOD}}")]:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(placeholder)

    document.add_heading("1. Introduction", level=1)
    document.add_paragraph("{{INTRODUCTION_TEXT}}")

    document.add_heading("2. Monthly Overview", level=1)
    document.add_paragraph("Here is a summary of key performance indicators for {{MONTH_PERIOD}}:")
    document.add_paragraph("{{TABLE_PLACEHOLDER}}")

    document.add_heading("3. Conclusion", level=1)
    document.add_paragraph("{{CONCLUSION_TEXT}}")

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "Generated automatically by {{AUTHOR_NAME}} - {{COMPANY_NAME}}"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    document.save(path)
    print(f"Created {path}")


def create_pdf(path, pages):
    """`pages` is a list of (heading, [lines]) tuples, one per page."""
    c = canvas.Canvas(path, pagesize=A4)
    width, height = A4
    for page_number, (heading, lines) in enumerate(pages, start=1):
        c.setFont("Helvetica-Bold", 20)
        c.drawString(2.5 * cm, height - 3 * cm, heading)
        c.setLineWidth(1)
        c.line(2.5 * cm, height - 3.4 * cm, width - 2.5 * cm, height - 3.4 * cm)

        c.setFont("Helvetica", 12)
        y = height - 4.5 * cm
        for line in lines:
            if line.startswith("## "):
                c.setFont("Helvetica-Bold", 14)
                y -= 0.3 * cm
                c.drawString(2.5 * cm, y, line[3:])
                c.setFont("Helvetica", 12)
            else:
                c.drawString(2.5 * cm, y, line)
            y -= 0.7 * cm

        c.setFont("Helvetica-Oblique", 9)
        c.drawCentredString(width / 2, 1.5 * cm, f"Page {page_number} of {len(pages)}")
        c.showPage()
    c.save()
    print(f"Created {path}")


def main():
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    os.makedirs(DOCUMENTS_DIR, exist_ok=True)

    create_word_template(os.path.join(TEMPLATES_DIR, "report_template.docx"))

    create_pdf(os.path.join(DOCUMENTS_DIR, "source_document.pdf"), [
        ("Enhanced Task Manager - Project Proposal", [
            "Project Name: Enhanced Task Manager",
            "Document ID: PROJ-PROP-011",
            "Version: 1.0",
            "",
            "## 1. Background",
            "Small teams often track their work in scattered spreadsheets and chats.",
            "The Enhanced Task Manager collects tasks in one place and automates",
            "the reporting work around them.",
            "",
            "## 2. Goals",
            "- Create, update and complete tasks with priorities and due dates",
            "- Generate a monthly Word report automatically from a template",
            "- Protect shared PDFs with a CONFIDENTIAL watermark",
            "- Merge project documents into a single PDF",
        ]),
        ("Enhanced Task Manager - Design", [
            "## 3. Architecture",
            "A configuration-driven Document Agent reads a JSON file that lists",
            "the tasks to run. Each task type is dispatched to a handler:",
            "- generate_word_report  -> WordTasks",
            "- apply_pdf_watermark   -> PDFTasks",
            "- merge_pdfs            -> PDFTasks",
            "- extract_word_text / extract_pdf_text",
            "",
            "## 4. Audit Logging",
            "Every task writes an entry (timestamp, status, message) to an",
            "audit report saved as JSON in documents/reports/.",
        ]),
    ])

    create_pdf(os.path.join(DOCUMENTS_DIR, "dummy_doc1.pdf"), [
        ("Project Documentation - Part A", [
            "Project Name: Agentic Document Automation System",
            "Document ID: PROJ-DOC-001",
            "",
            "## Chapter 1: Introduction",
            "The system lets users define document processing pipelines",
            "through simple JSON configuration files.",
            "",
            "Key Features:",
            "- Template-based document generation",
            "- Dynamic content insertion",
            "- Advanced PDF manipulation",
            "- Comprehensive audit logging",
            "",
            "--- END OF PART A ---",
        ]),
    ])

    create_pdf(os.path.join(DOCUMENTS_DIR, "dummy_doc2.pdf"), [
        ("Project Documentation - Part B", [
            "Project Name: Agentic Document Automation System",
            "Document ID: PROJ-DOC-002",
            "",
            "## Chapter 2: Implementation",
            "- python-docx fills Word templates and builds tables",
            "- reportlab draws the watermark text",
            "- pypdf merges pages and stamps the watermark onto each page",
            "",
            "## Chapter 3: Testing",
            "Run main.py and compare the files in documents/ with the",
            "expected results described in README.md.",
            "",
            "--- END OF PART B ---",
        ]),
    ])


if __name__ == "__main__":
    main()
