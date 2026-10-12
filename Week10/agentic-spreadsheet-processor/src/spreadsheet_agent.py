# agentic-spreadsheet-processor/src/spreadsheet_agent.py
import datetime
import json
from copy import copy

import openpyxl
from openpyxl.styles import Font, PatternFill, Color
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, IconSetRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter, range_boundaries

from .utils import setup_logging # Use utils for consistent logging setup

logger = setup_logging(__name__)

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="333333", end_color="333333", fill_type="solid") # Dark Gray

class SpreadsheetAgent:
    """
    An agent that performs a series of operations on an Excel workbook
    based on a provided configuration.
    """
    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []
        self.start_time = datetime.datetime.now()

        # Task dispatcher: maps the task 'type' in the config to its handler method
        self.task_handlers = {
            "copy_data": self._process_copy_data_task,
            "calculate_column": self._process_calculate_column_task,
            "conditional_format": self._process_conditional_format_task,
            "chart": self._process_chart_task,
            "summary": self._process_summary_task,
        }
        logger.info("Spreadsheet Agent initialized.")

    def _log_audit(self, status, message, task_type="N/A", details=None):
        """Adds an entry to the audit log."""
        log_entry = {
            "timestamp": datetime.datetime.now().isoformat(timespec="seconds"),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details if details else {}
        }
        self.audit_log.append(log_entry)
        if status == "ERROR" or status == "CRITICAL":
            logger.error(f"AUDIT LOG - {status}: {message} | Task: {task_type} | Details: {details}")
        elif status == "WARNING":
            logger.warning(f"AUDIT LOG - {status}: {message} | Task: {task_type}")
        else:
            logger.info(f"AUDIT LOG - {status}: {message} | Task: {task_type}")

    def _load_workbook(self):
        """Loads the input workbook."""
        input_file = self.config["input_file"]
        try:
            self.workbook = openpyxl.load_workbook(input_file)
            self._log_audit("SUCCESS", f"Workbook '{input_file}' loaded.", "load_workbook")
            return True
        except FileNotFoundError:
            self._log_audit("ERROR", f"Input file not found: {input_file}", "load_workbook")
            return False
        except Exception as e:
            self._log_audit("ERROR", f"Failed to load workbook '{input_file}': {e}", "load_workbook")
            return False

    def _get_sheet(self, sheet_name, create=True):
        """Safely gets a worksheet from the workbook, creating it if allowed."""
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        if not create:
            return None
        self._log_audit("INFO", f"Sheet '{sheet_name}' not found. Creating new sheet.", "get_sheet")
        return self.workbook.create_sheet(sheet_name)

    @staticmethod
    def _last_data_row(sheet, column="A"):
        """Returns the last row that has a value in `column` (so empty trailing rows are ignored)."""
        for row_num in range(sheet.max_row, 0, -1):
            if sheet[f"{column}{row_num}"].value not in (None, ""):
                return row_num
        return 1

    @staticmethod
    def _style_header_cell(cell):
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    def _process_copy_data_task(self, task):
        """Copies data (values, number formats and fonts) from a source sheet to a destination sheet."""
        src_sheet_name = task.get("source_sheet")
        dest_sheet_name = task.get("dest_sheet")

        src_sheet = self._get_sheet(src_sheet_name, create=False)
        if src_sheet is None:
            self._log_audit("ERROR", f"Source sheet '{src_sheet_name}' not found for copying.", "copy_data", task)
            return

        dest_sheet = self._get_sheet(dest_sheet_name)
        for row in src_sheet.iter_rows():
            for cell in row:
                dest_cell = dest_sheet[cell.coordinate]
                dest_cell.value = cell.value
                dest_cell.number_format = cell.number_format
                dest_cell.font = copy(cell.font)

        # Header row styling and column widths
        for cell in dest_sheet[1]:
            if cell.value is not None:
                self._style_header_cell(cell)
        for letter, dimension in src_sheet.column_dimensions.items():
            dest_sheet.column_dimensions[letter].width = dimension.width
        dest_sheet.freeze_panes = "A2"

        self._log_audit("SUCCESS", f"Data copied from '{src_sheet_name}' to '{dest_sheet_name}' ({src_sheet.max_row} rows).", "copy_data")

    def _process_calculate_column_task(self, task):
        """Fills a column with an Excel formula; '{row}' in the formula is replaced by the row number."""
        sheet_name = task.get("sheet")
        target_column_letter = task.get("target_column")
        start_row = task.get("start_row", 2) # Assume header in row 1
        formula = task.get("formula") # e.g., "=B{row}*C{row}"

        sheet = self._get_sheet(sheet_name)

        header_cell = sheet[f"{target_column_letter}1"]
        header_cell.value = task.get("header", "Calculated Column")
        self._style_header_cell(header_cell)
        sheet.column_dimensions[target_column_letter].width = task.get("column_width", 16)

        end_row = self._last_data_row(sheet)
        failed = 0
        for row_num in range(start_row, end_row + 1):
            cell_coord = f"{target_column_letter}{row_num}"
            try:
                sheet[cell_coord] = formula.replace("{row}", str(row_num))
                if task.get("number_format"):
                    sheet[cell_coord].number_format = task["number_format"]
            except Exception as e:
                failed += 1
                self._log_audit("ERROR", f"Failed to apply formula to {cell_coord}: {e}", "calculate_column", task)

        status = "SUCCESS" if failed == 0 else "WARNING"
        self._log_audit(status, f"Column '{target_column_letter}' ({header_cell.value}) calculated in '{sheet_name}' "
                                f"for rows {start_row}-{end_row}.", "calculate_column")

    def _process_conditional_format_task(self, task):
        """Applies conditional formatting to a range."""
        sheet_name = task.get("sheet")
        data_range = task.get("range")
        rule_type = task.get("rule_type")

        sheet = self._get_sheet(sheet_name)

        try:
            if rule_type == "color_scale":
                # 3-color scale: min -> 50th percentile -> max
                rule = ColorScaleRule(start_type='min', start_color=Color(task.get("min_color", "FF0000")),
                                      mid_type='percentile', mid_value=50, mid_color=Color(task.get("mid_color", "FFFF00")),
                                      end_type='max', end_color=Color(task.get("max_color", "00FF00")))
            elif rule_type == "data_bar":
                rule = DataBarRule(start_type='min', end_type='max', color=Color(task.get("color", "0070C0")))
            elif rule_type == "icon_set":
                rule = IconSetRule(task.get("icon_set", '3TrafficLights1'), task.get("value_type", 'percent'),
                                   task.get("values", [0, 33, 67]))
            elif rule_type == "expression":
                # Highlight cells where the formula is TRUE. Excel stores the formula without the leading '='.
                formula = task.get("formula", "").lstrip("=")
                dxf = DifferentialStyle(font=Font(color=Color(task.get("font_color", "9C0006"))),
                                        fill=PatternFill(bgColor=Color(task.get("fill_color", "FFC7CE"))))
                rule = Rule(type='expression', dxf=dxf, formula=[formula])
            else:
                self._log_audit("ERROR", f"Unsupported conditional format rule_type: {rule_type}", "conditional_format", task)
                return

            sheet.conditional_formatting.add(data_range, rule)
            self._log_audit("SUCCESS", f"Conditional formatting '{rule_type}' applied to '{data_range}' in '{sheet_name}'.", "conditional_format")

        except Exception as e:
            self._log_audit("ERROR", f"Failed to apply conditional formatting to '{data_range}': {e}", "conditional_format", task)

    def _process_chart_task(self, task):
        """Creates a bar/line/pie chart from a data range (header row included) and category labels."""
        sheet_name = task.get("sheet")
        chart_type = task.get("chart_type")
        data_range = task.get("data_range") # e.g., "D1:D10" (row 1 = series title)
        category_range = task.get("category_range") # e.g., "A2:A10" (labels, no header)
        title = task.get("title", "Chart Title")
        top_left_cell = task.get("top_left_cell", "G2") # Where to place the chart

        sheet = self._get_sheet(sheet_name)

        chart_classes = {"bar": BarChart, "line": LineChart, "pie": PieChart}
        if chart_type not in chart_classes:
            self._log_audit("ERROR", f"Unsupported chart type: {chart_type}", "chart", task)
            return

        try:
            chart = chart_classes[chart_type]()

            min_col, min_row, max_col, max_row = range_boundaries(data_range)
            data = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.add_data(data, titles_from_data=True) # First row of data_range is the series title

            if category_range:
                min_col, min_row, max_col, max_row = range_boundaries(category_range)
                chart.set_categories(Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row))

            chart.title = title
            if chart_type != "pie":
                chart.x_axis.title = task.get("x_axis_title", "")
                chart.y_axis.title = task.get("y_axis_title", "")
                # Newer Excel versions hide axes unless 'delete' is explicitly False
                chart.x_axis.delete = False
                chart.y_axis.delete = False
            chart.width = task.get("width", 18)
            chart.height = task.get("height", 9)

            sheet.add_chart(chart, top_left_cell)
            self._log_audit("SUCCESS", f"Chart '{title}' of type '{chart_type}' added to '{sheet_name}' at {top_left_cell}.", "chart")

        except Exception as e:
            self._log_audit("ERROR", f"Failed to create chart '{title}': {e}", "chart", task)

    def _process_summary_task(self, task):
        """Writes a 'Metric | Value' table of aggregate formulas (SUM, AVERAGE, ...) to a summary sheet."""
        sheet_name = task.get("sheet")
        sheet = self._get_sheet(sheet_name)

        sheet["A1"], sheet["B1"] = "Metric", "Value"
        self._style_header_cell(sheet["A1"])
        self._style_header_cell(sheet["B1"])

        for row_num, metric in enumerate(task.get("metrics", []), start=2):
            sheet[f"A{row_num}"] = metric.get("label", f"Metric {row_num - 1}")
            sheet[f"B{row_num}"] = metric.get("formula")
            if metric.get("number_format"):
                sheet[f"B{row_num}"].number_format = metric["number_format"]

        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 18
        self._log_audit("SUCCESS", f"Summary with {len(task.get('metrics', []))} metrics written to '{sheet_name}'.", "summary")

    def _add_audit_log_sheet(self):
        """Adds a sheet to the workbook containing the audit log."""
        if not self.audit_log:
            return

        audit_sheet = self.workbook.create_sheet("Audit Log", 0) # Create as first sheet
        audit_sheet.append(["Timestamp", "Status", "Task Type", "Message", "Details"])
        for cell in audit_sheet[1]:
            self._style_header_cell(cell)

        status_fills = {
            "SUCCESS": "C6EFCE", # green
            "INFO": "DDEBF7",    # blue
            "WARNING": "FFEB9C", # yellow
            "ERROR": "FFC7CE",   # red
            "CRITICAL": "FFC7CE",
        }
        for entry in self.audit_log:
            audit_sheet.append([
                entry["timestamp"],
                entry["status"],
                entry["task_type"],
                entry["message"],
                json.dumps(entry["details"], ensure_ascii=False) if entry["details"] else "" # Details as JSON string
            ])
            fill_color = status_fills.get(entry["status"])
            if fill_color:
                audit_sheet.cell(row=audit_sheet.max_row, column=2).fill = PatternFill(
                    start_color=fill_color, end_color=fill_color, fill_type="solid")

        # Auto-adjust column widths (capped so the Details column stays readable)
        for col_num in range(1, 6):
            column_letter = get_column_letter(col_num)
            max_length = max((len(str(cell.value)) for cell in audit_sheet[column_letter] if cell.value), default=0)
            audit_sheet.column_dimensions[column_letter].width = min(max_length + 2, 80)
        audit_sheet.freeze_panes = "A2"

        logger.info("Audit log sheet created.")

    def run(self):
        """Executes the defined tasks from the configuration."""
        if not self._load_workbook():
            self._log_audit("CRITICAL", "Agent could not start due to workbook loading failure.", "run")
            return False

        try:
            # Process each task sequentially
            for i, task in enumerate(self.config["tasks"], start=1):
                task_type = task.get("type")
                logger.info(f"Executing Task {i}: Type='{task_type}' Sheet='{task.get('sheet', task.get('dest_sheet'))}'")

                handler = self.task_handlers.get(task_type)
                if handler is None:
                    self._log_audit("WARNING", f"Unsupported task type encountered: {task_type}", "task_execution", task)
                    continue
                try:
                    handler(task)
                except Exception as e:
                    # One failing task must not stop the rest of the workflow
                    self._log_audit("ERROR", f"Task {i} ({task_type}) failed: {e}", task_type, task)

            duration = (datetime.datetime.now() - self.start_time).total_seconds()
            output_file = self.config["output_file"]
            self._log_audit("INFO", f"All {len(self.config['tasks'])} tasks processed in {duration:.2f} seconds. "
                                    f"Saving workbook to '{output_file}'.", "run_summary")

            # Add final audit log sheet, then save
            self._add_audit_log_sheet()
            # openpyxl does not calculate formulas, so ask Excel to recalculate everything when the file opens
            self.workbook.calculation.fullCalcOnLoad = True
            self.workbook.save(output_file)
            logger.info(f"Final workbook saved to: {output_file}")
            return True

        except Exception as e:
            logger.critical(f"An unhandled error occurred during agent execution: {e}", exc_info=True)
            return False
        finally:
            duration = (datetime.datetime.now() - self.start_time).total_seconds()
            logger.info(f"Agent finished. Total time: {duration:.2f} seconds.")
