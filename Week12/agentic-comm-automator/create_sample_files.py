# agentic-comm-automator/create_sample_files.py
# Creates the sample attachments used by the email tasks:
#   data/reports/sales_report.pdf   and   data/reports/summary_data.xlsx
import os

import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REPORTS_DIR = os.path.join(BASE_DIR, "data", "reports")

SALES = [
    ("Laptop", 2, 1200.50), ("Monitor", 4, 300.00), ("Projector", 3, 850.00),
    ("Webcam", 25, 45.00), ("Headset", 30, 99.99), ("SSD", 20, 180.00),
]

def create_pdf(path):
    c = canvas.Canvas(path, pagesize=A4)
    width, height = A4
    c.setFont("Helvetica-Bold", 20)
    c.drawString(2.5 * cm, height - 3 * cm, "Daily Sales Report")
    c.setFont("Helvetica", 11)
    c.drawString(2.5 * cm, height - 3.8 * cm, "Generated for the Agentic Communication Automator lab")

    y = height - 5.5 * cm
    c.setFont("Helvetica-Bold", 12)
    for x, header in zip([2.5, 8, 11, 14.5], ["Product", "Qty", "Unit Price", "Revenue"]):
        c.drawString(x * cm, y, header)
    c.line(2.5 * cm, y - 0.3 * cm, 18 * cm, y - 0.3 * cm)
    c.setFont("Helvetica", 12)
    total = 0
    for product, qty, price in SALES:
        y -= 0.8 * cm
        revenue = qty * price
        total += revenue
        for x, value in zip([2.5, 8, 11, 14.5], [product, str(qty), f"{price:,.2f}", f"{revenue:,.2f}"]):
            c.drawString(x * cm, y, value)
    y -= 1.0 * cm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(2.5 * cm, y, "Total")
    c.drawString(14.5 * cm, y, f"{total:,.2f}")
    c.save()
    print(f"Created {path}")

def create_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Product", "Quantity", "Unit Price", "Revenue"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    for row, (product, qty, price) in enumerate(SALES, start=2):
        ws.append([product, qty, price, f"=B{row}*C{row}"])
    ws.append(["Total", f"=SUM(B2:B{len(SALES) + 1})", None, f"=SUM(D2:D{len(SALES) + 1})"])
    for column, width in zip("ABCD", [14, 10, 12, 12]):
        ws.column_dimensions[column].width = width
    wb.save(path)
    print(f"Created {path}")

if __name__ == "__main__":
    os.makedirs(REPORTS_DIR, exist_ok=True)
    create_pdf(os.path.join(REPORTS_DIR, "sales_report.pdf"))
    create_xlsx(os.path.join(REPORTS_DIR, "summary_data.xlsx"))
