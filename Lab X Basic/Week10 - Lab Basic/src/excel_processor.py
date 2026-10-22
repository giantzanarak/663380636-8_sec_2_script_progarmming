# spreadsheet-automator/src/excel_processor.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ["Product Name", "Quantity", "Unit Price", "Total Price"]

# Styles used in the report
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")   # green
TOTAL_FONT = Font(bold=True, color="000000")
TOTAL_FILL = PatternFill(start_color="FFEB3B", end_color="FFEB3B", fill_type="solid")    # yellow
THIN = Side(style="thin")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CURRENCY_FORMAT = "#,##0.00"


class ExcelProcessor:
    """
    Common Excel processing functions using openpyxl.
    """

    def create_workbook(self):
        """Creates a new, empty Excel workbook (the default 'Sheet' is removed)."""
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        print("New workbook created.")
        return wb

    def load_workbook(self, filename):
        """Loads an existing Excel workbook. Returns None if it cannot be opened."""
        try:
            wb = openpyxl.load_workbook(filename, data_only=True)  # read values, not formulas
            print(f"Workbook '{filename}' loaded successfully.")
            return wb
        except FileNotFoundError:
            print(f"Error: Workbook '{filename}' not found.")
        except Exception as e:
            print(f"Error loading workbook '{filename}': {e}")
        return None

    def save_workbook(self, workbook, filename):
        """Saves a workbook. Returns True on success."""
        try:
            workbook.save(filename)
            print(f"Workbook saved to '{filename}'.")
            return True
        except PermissionError:
            print(f"Error: cannot write '{filename}'. Close it in Excel and try again.")
        except Exception as e:
            print(f"Error saving workbook to '{filename}': {e}")
        return False

    @staticmethod
    def _read_sales_rows(sheet):
        """
        Reads (product, quantity, unit_price) from the sheet, skipping the header row,
        empty rows and rows whose numbers are invalid.
        """
        rows = []
        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not any(value not in (None, "") for value in row):
                continue  # empty row
            product_name = row[0] if len(row) > 0 else None
            try:
                quantity = int(row[1])
                unit_price = float(row[2])
            except (ValueError, TypeError, IndexError):
                print(f"Warning: Skipping row {row_index} due to invalid number format: {row}")
                continue
            if not product_name:
                print(f"Warning: Skipping row {row_index} because the product name is empty: {row}")
                continue
            rows.append((str(product_name).strip(), quantity, unit_price))
        return rows

    @staticmethod
    def _autofit_columns(sheet):
        """Sets each column width from its longest value (merged label cells are ignored)."""
        for col_num in range(1, len(HEADERS) + 1):
            letter = get_column_letter(col_num)
            longest = 0
            for cell in sheet[letter]:
                if cell.value is None:
                    continue
                text = f"{cell.value:,.2f}" if isinstance(cell.value, float) else str(cell.value)
                if cell.value != "Grand Total:":
                    longest = max(longest, len(text))
            sheet.column_dimensions[letter].width = longest + 4

    def process_sales_data(self, input_filename, output_filename):
        """
        Loads sales data, calculates the total price of each item and the grand total,
        and writes a formatted report to a new Excel file.
        Input columns (row 1 = header): Product Name, Quantity, Unit Price.
        """
        input_wb = self.load_workbook(input_filename)
        if not input_wb:
            return False

        sales = self._read_sales_rows(input_wb.active)
        if not sales:
            print("Error: no valid sales rows found in the input file.")
            return False

        output_wb = self.create_workbook()
        sheet = output_wb.create_sheet("Sales Report", 0)

        # Header row
        for col_num, header in enumerate(HEADERS, start=1):
            cell = sheet.cell(row=1, column=col_num, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows
        grand_total = 0.0
        for output_row, (product_name, quantity, unit_price) in enumerate(sales, start=2):
            total_price = round(quantity * unit_price, 2)
            grand_total += total_price
            values = [product_name, quantity, unit_price, total_price]
            for col_num, value in enumerate(values, start=1):
                cell = sheet.cell(row=output_row, column=col_num, value=value)
                cell.border = THIN_BORDER
                if col_num in (3, 4):
                    cell.number_format = CURRENCY_FORMAT

        # Grand total row: label merged across A:C, value in D
        total_row = len(sales) + 2
        sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
        for col_num in range(1, 5):
            cell = sheet.cell(row=total_row, column=col_num)
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = THIN_BORDER
        label = sheet.cell(row=total_row, column=1, value="Grand Total:")
        label.alignment = Alignment(horizontal="right", vertical="center")
        value = sheet.cell(row=total_row, column=4, value=round(grand_total, 2))
        value.number_format = CURRENCY_FORMAT

        sheet.freeze_panes = "A2"
        self._autofit_columns(sheet)

        print(f"Processed {len(sales)} product(s). Grand total: {grand_total:,.2f}")
        return self.save_workbook(output_wb, output_filename)
