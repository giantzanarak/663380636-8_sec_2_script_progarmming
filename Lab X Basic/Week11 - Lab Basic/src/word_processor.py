# document-automator/src/word_processor.py
import os

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from utils import setup_logging

logger = setup_logging(__name__)

ALIGNMENTS = {
    'left': WD_ALIGN_PARAGRAPH.LEFT,
    'center': WD_ALIGN_PARAGRAPH.CENTER,
    'right': WD_ALIGN_PARAGRAPH.RIGHT,
    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
}

class WordProcessor:
    """
    Handles operations on Word documents using python-docx.
    """
    def __init__(self):
        logger.info("WordProcessor initialized.")

    def create_document(self):
        """Creates a new blank Word document."""
        return docx.Document()

    def add_heading(self, document, text, level=1):
        """Adds a heading (level 0 = title, 1-9 = heading levels)."""
        document.add_heading(text, level=level)

    def add_paragraph(self, document, text, style=None, bold=False, italic=False, font_size=None, align='left'):
        """Adds a paragraph; formatting is applied to the whole paragraph."""
        paragraph = document.add_paragraph(style=style)
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        if font_size:
            run.font.size = Pt(font_size)
        paragraph.alignment = ALIGNMENTS.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        return paragraph

    def add_table(self, document, data, header_bold=True, col_widths=None):
        """
        Adds a table. `data` is a list of rows; the first row is the header.
        `col_widths` is an optional list of widths in inches.
        """
        if not data or not data[0]:
            logger.warning("Table data is empty or malformed. Skipping table creation.")
            return None

        cols = len(data[0])
        table = document.add_table(rows=len(data), cols=cols)
        table.style = 'Table Grid'

        for r_idx, row_data in enumerate(data):
            cells = table.rows[r_idx].cells
            for c_idx, value in enumerate(row_data[:cols]):
                cells[c_idx].text = str(value)
                if r_idx == 0 and header_bold:
                    for run in cells[c_idx].paragraphs[0].runs:
                        run.bold = True

        if col_widths and len(col_widths) == cols:
            # Word uses the cell widths, so set them on every cell of each column
            for c_idx, width in enumerate(col_widths):
                for cell in table.columns[c_idx].cells:
                    cell.width = Inches(width)

        logger.info(f"Added table with {len(data)} rows and {cols} columns.")
        return table

    def add_picture(self, document, image_path, width_inches=None):
        """Adds an image with optional width scaling. Returns True on success."""
        if not os.path.exists(image_path):
            logger.error(f"Image file not found: {image_path}")
            return False
        try:
            document.add_picture(image_path, width=Inches(width_inches) if width_inches else None)
            return True
        except Exception as e:
            logger.error(f"Error adding picture '{image_path}': {e}")
            return False

    def save_document(self, document, output_path):
        """Saves the Word document. Returns True on success."""
        try:
            document.save(output_path)
            logger.info(f"Word document saved to '{os.path.basename(output_path)}'.")
            return True
        except PermissionError:
            logger.error(f"Cannot write '{output_path}'. Close it in Word and try again.")
        except Exception as e:
            logger.error(f"Error saving Word document to '{output_path}': {e}")
        return False
